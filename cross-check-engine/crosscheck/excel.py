"""Nhập hồ sơ dự án từ Excel (CR-003, đặc tả docs/spec/SPEC-CR-003.md).

    python3 -m crosscheck.excel mau ho_so_mau.xlsx            # tạo file mẫu trống
    python3 -m crosscheck.excel tu-yaml ho_so.yaml ho_so.xlsx  # chuyển hồ sơ YAML sang Excel

Đọc: `load_project_xlsx(path, rules, cfg)` trả về hồ sơ cùng cấu trúc với YAML; lỗi nhập liệu được gom hết
thành `ExcelInputError` (mỗi lỗi ghi Sheet!Ô và cách sửa). Engine chỉ đọc file, không bao giờ ghi vào file
của người dùng. Cần `openpyxl` (chỉ khi dùng Excel); nên cài thêm `defusedxml` - openpyxl tự dùng nó để chặn
XML độc hại trong file đến từ bên ngoài.
"""
from __future__ import annotations

import datetime as dt
import re
import sys
import zipfile
from pathlib import Path

from .engine import RuleError

MAX_BYTES = 20 * 1024 * 1024
STATUS = {"not_started": "Chưa bắt đầu", "in_progress": "Đang thực hiện", "completed": "Hoàn thành",
          "not_applicable": "Không áp dụng"}
ATTRS = [("requires_permit", "Thuộc diện cấp GPXD"), ("requires_design_verification", "Thuộc diện thẩm tra thiết kế"),
         ("requires_safety_certification", "Thuộc diện chứng nhận an toàn"), ("von_dau_tu_cong", "Vốn đầu tư công")]
TRINHTU = [("step", "Bước *"), ("status", "Trạng thái *"), ("start_date", "Ngày bắt đầu"),
           ("date", "Ngày hoàn thành"), ("evidence", "Chứng cứ"), ("na_basis", "Căn cứ không áp dụng")]
HOPDONG = [("id", "Mã hợp đồng *"), ("ngay_ky", "Ngày ký *")]
VANBAN = [("id", "Mã văn bản *"), ("loai", "Loại văn bản"), ("nhom", "Nhóm"), ("ngay_ky", "Ngày ký"),
          ("ngay_nop", "Ngày nộp"), ("ngay_phat_hanh_hsmt", "Ngày phát hành HSMT"), ("hop_dong", "Mã hợp đồng"),
          ("thay_the", "Thay thế (mã, cách nhau dấu phẩy)"), ("ten_du_an", "Tên dự án"),
          ("tong_muc_dau_tu", "TMĐT (đồng)"), ("von_duoc_giao", "Vốn được giao (đồng)"),
          ("dien_tich_san", "Diện tích sàn (m²)"), ("chieu_cao", "Chiều cao (m)"), ("so_tang", "Số tầng"),
          ("tg_tu", "Thời gian thực hiện từ"), ("tg_den", "Thời gian thực hiện đến")]
CANCU = [("doc", "Mã văn bản *"), ("van_ban", "Văn bản viện dẫn *"), ("dieu_khoan", "Điều/khoản")]
INFO_TEXT, INFO_NUM = ("ten_du_an",), ("tong_muc_dau_tu", "von_duoc_giao", "dien_tich_san", "chieu_cao", "so_tang")
MONEY_COLS = {"tong_muc_dau_tu", "von_duoc_giao"}
WIDE = {"step": 52, "nhom": 34, "loai": 34, "van_ban": 40, "ten_du_an": 28, "na_basis": 30}  # nhãn dài không bị cắt
DATE_COLS = {"start_date", "date", "ngay_ky", "ngay_nop", "ngay_phat_hanh_hsmt", "tg_tu", "tg_den"}
SHEETS = ["HuongDan", "DuAn", "TrinhTu", "HopDong", "VanBan", "CanCu"]


class ExcelInputError(RuleError):  # CLI bắt RuleError → in danh sách lỗi, mã thoát 2
    def __init__(self, path, errors):
        self.errors = errors
        super().__init__(f"file {Path(path).name} có {len(errors)} lỗi nhập liệu, chưa chạy đối soát.\n  "
                         + "\n  ".join(errors))


def _openpyxl():
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover
        raise SystemExit("Cần thư viện openpyxl để dùng Excel: pip install openpyxl") from e
    return openpyxl


def _norm_header(h) -> str:
    return " ".join(str(h or "").replace("*", "").replace("▼", "").split()).lower()


def _key(s) -> str:
    return " ".join(str(s).split()).lower()


# ---------------------------------------------------------------- đọc

def load_project_xlsx(path, rules: dict, cfg: dict) -> dict:
    path = Path(path)
    errors: list[str] = []
    if path.suffix.lower() != ".xlsx":
        raise ExcelInputError(path, [f"{path.name}: chỉ nhận file .xlsx (Excel 2007 trở lên)."])
    if path.stat().st_size > MAX_BYTES:
        raise ExcelInputError(path, [f"{path.name}: file lớn hơn 20 MB."])
    if not zipfile.is_zipfile(path):
        raise ExcelInputError(path, [f"{path.name}: không phải file .xlsx hợp lệ (có thể bị hỏng hoặc đổi đuôi)."])
    opx = _openpyxl()
    try:
        wb = opx.load_workbook(path, data_only=True)        # giá trị đã tính (B6)
        wbf = opx.load_workbook(path, data_only=False)      # để nhận ra ô công thức
    except Exception as e:  # noqa: BLE001 - mọi lỗi đọc file đều trả về người dùng một thông điệp
        raise ExcelInputError(path, [f"{path.name}: không đọc được file ({type(e).__name__})."]) from e

    step_by = {}
    for sid, r in rules.items():
        step_by[_key(sid)] = sid
        step_by[_key(r["name"])] = sid
    group_by = {}
    for g, label in cfg["groups"].items():
        group_by[_key(g)] = g
        group_by[_key(label)] = g
    status_by = {**{_key(k): k for k in STATUS}, **{_key(v): k for k, v in STATUS.items()}}

    def cellref(ws, row, col):
        return f"{ws.title}!{opx.utils.get_column_letter(col)}{row}"

    def value(ws, row, col):
        v = ws.cell(row=row, column=col).value
        fv = wbf[ws.title].cell(row=row, column=col).value
        if v is None and isinstance(fv, str) and fv.startswith("="):
            errors.append(f"{cellref(ws, row, col)}  Ô công thức chưa có giá trị đã tính. Mở file bằng Excel, lưu lại, "
                          "hoặc nhập giá trị trực tiếp.")
            return None
        if isinstance(v, str):
            v = v.strip()
            return v or None
        return v

    def as_date(v, ref, label):
        if v is None:
            return None
        if isinstance(v, dt.datetime):
            return v.date()
        if isinstance(v, dt.date):
            return v
        s = str(v)
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        errors.append(f"{ref}  {label} \"{s}\" không phải ngày hợp lệ. Dùng dd/mm/yyyy.")
        return None

    def as_number(v, ref, label):
        if v is None:
            return None
        if isinstance(v, bool):
            errors.append(f"{ref}  {label} phải là số.")
            return None
        if isinstance(v, (int, float)):
            return int(v) if isinstance(v, float) and v.is_integer() else v
        s = str(v).replace(" ", "")
        if re.fullmatch(r"-?\d{1,3}(\.\d{3})*(,\d+)?|-?\d+(,\d+)?", s):
            s = s.replace(".", "").replace(",", ".")
            n = float(s)
            return int(n) if n.is_integer() and "." not in s else n
        errors.append(f"{ref}  {label} \"{v}\" không phải số. Nhập số, hoặc chữ theo kiểu 120.000.000.000 / 24,5.")
        return None

    def table(name, cols, required=True):
        """Đọc sheet dạng bảng: [(dòng, {khóa: giá trị})]. Cột tìm theo tiêu đề, không theo vị trí."""
        if name not in wb.sheetnames:
            if required:
                errors.append(f"{name}  Thiếu sheet bắt buộc \"{name}\". Dùng file mẫu: python3 -m crosscheck.excel mau.")
            return None, []
        ws = wb[name]
        headers = {_norm_header(ws.cell(row=1, column=c).value): c for c in range(1, ws.max_column + 1)}
        idx = {}
        for k, label in cols:
            c = headers.get(_norm_header(label))
            if c is None:
                if label.endswith("*"):
                    errors.append(f"{name}!1  Thiếu cột bắt buộc \"{label.rstrip(' *')}\".")
                continue
            idx[k] = c
        rows = []
        for r in range(2, ws.max_row + 1):
            vals = {k: value(ws, r, c) for k, c in idx.items()}
            if all(v is None for v in vals.values()):
                continue
            for k, label in cols:
                if label.endswith("*") and k in idx and vals.get(k) is None:
                    errors.append(f"{cellref(ws, r, idx[k])}  Ô bắt buộc \"{label.rstrip(' *')}\" đang trống.")
            rows.append((r, vals, ws, idx))
        return ws, rows

    def ref(item, k):
        r, _vals, ws, idx = item
        return cellref(ws, r, idx[k])

    # DuAn
    project, attributes = {}, {}
    if "DuAn" not in wb.sheetnames:
        errors.append("DuAn  Thiếu sheet bắt buộc \"DuAn\". Dùng file mẫu: python3 -m crosscheck.excel mau.")
    else:
        ws = wb["DuAn"]
        rows = {_norm_header(ws.cell(row=r, column=1).value): r for r in range(2, ws.max_row + 1)}
        for key, label in (("id", "Mã dự án"), ("name", "Tên dự án")):
            r = rows.get(_norm_header(label))
            v = value(ws, r, 2) if r else None
            if v is None:
                errors.append(f"DuAn!B{r or '?'}  Ô bắt buộc \"{label}\" đang trống.")
            else:
                project[key] = str(v)
        for key, label in ATTRS:
            r = rows.get(_norm_header(label))
            v = value(ws, r, 2) if r else None
            if v is None:
                continue
            k = _key(v)
            if k in ("có", "co", "true"):          # nhãn hoặc mã (B2) - không đoán "x", "1"
                attributes[key] = True
            elif k in ("không", "khong", "false"):
                attributes[key] = False
            else:
                errors.append(f"DuAn!B{r}  \"{v}\" không hợp lệ cho \"{label}\". Chọn: Có, Không, hoặc để trống.")

    # TrinhTu
    steps = {}
    _ws, rows = table("TrinhTu", TRINHTU, required=False)
    for item in rows:
        _r, v, _w, _i = item
        if v.get("step") is None:
            continue
        sid = step_by.get(_key(v["step"]))
        if sid is None:
            errors.append(f"{ref(item, 'step')}  Bước \"{v['step']}\" không có trong danh mục bước (xem sheet HuongDan).")
            continue
        if sid in steps:
            errors.append(f"{ref(item, 'step')}  Bước \"{v['step']}\" khai báo trùng.")
            continue
        rec = {}
        if v.get("status") is not None:
            st = status_by.get(_key(v["status"]))
            if st is None:
                errors.append(f"{ref(item, 'status')}  Trạng thái \"{v['status']}\" không hợp lệ. Chọn: "
                              + ", ".join(STATUS.values()) + ".")
            else:
                rec["status"] = st
        for k, label in (("start_date", "Ngày bắt đầu"), ("date", "Ngày hoàn thành")):
            if k in v and v[k] is not None:
                d = as_date(v[k], ref(item, k), label)
                if d:
                    rec[k] = d
        for k in ("evidence", "na_basis"):
            if v.get(k) is not None:
                rec[k] = str(v[k])
        steps[sid] = rec

    # HopDong
    contracts = {}
    _ws, rows = table("HopDong", HOPDONG, required=False)
    for item in rows:
        v = item[1]
        if v.get("id") is None:
            continue
        cid = str(v["id"])
        if cid in contracts:
            errors.append(f"{ref(item, 'id')}  Mã hợp đồng \"{cid}\" trùng.")
            continue
        contracts[cid] = {"ngay_ky": as_date(v.get("ngay_ky"), ref(item, "ngay_ky"), "Ngày ký")} if "ngay_ky" in v else {}

    # VanBan
    documents, doc_ids = [], {}
    _ws, rows = table("VanBan", VANBAN, required=True)
    for item in rows:
        v = item[1]
        if v.get("id") is None:
            continue
        did = str(v["id"])
        if did in doc_ids:
            errors.append(f"{ref(item, 'id')}  Mã văn bản \"{did}\" trùng.")
            continue
        d = {"id": did}
        if v.get("loai") is not None:
            d["loai"] = str(v["loai"])
        if v.get("nhom") is not None:
            g = group_by.get(_key(v["nhom"]))
            if g is None:
                errors.append(f"{ref(item, 'nhom')}  Nhóm \"{v['nhom']}\" không có trong danh mục nhóm (xem sheet HuongDan).")
            else:
                d["nhom"] = g
        for k, label in (("ngay_ky", "Ngày ký"), ("ngay_nop", "Ngày nộp"), ("ngay_phat_hanh_hsmt", "Ngày phát hành HSMT")):
            if v.get(k) is not None:
                dd = as_date(v[k], ref(item, k), label)
                if dd:
                    d[k] = dd
        if v.get("hop_dong") is not None:
            d["hop_dong"] = str(v["hop_dong"])
            if d["hop_dong"] not in contracts:
                errors.append(f"{ref(item, 'hop_dong')}  Mã hợp đồng \"{d['hop_dong']}\" không có trong sheet HopDong.")
        if v.get("thay_the") is not None:
            d["thay_the"] = [x.strip() for x in str(v["thay_the"]).split(",") if x.strip()]
        info = {}
        for k in INFO_TEXT:
            if v.get(k) is not None:
                info[k] = str(v[k])
        for k in INFO_NUM:
            if v.get(k) is not None:
                n = as_number(v[k], ref(item, k), dict(VANBAN)[k])
                if n is not None:
                    info[k] = n
        tg = {}
        for k, out, label in (("tg_tu", "tu", "Thời gian thực hiện từ"), ("tg_den", "den", "Thời gian thực hiện đến")):
            if v.get(k) is not None:
                dd = as_date(v[k], ref(item, k), label)
                if dd:
                    tg[out] = dd
        if tg:
            info["thoi_gian_thuc_hien"] = tg
        if info:
            d["thong_tin"] = info
        doc_ids[did] = (item, d)
        documents.append(d)
    for did, (item, d) in doc_ids.items():
        for old in d.get("thay_the") or []:
            if old not in doc_ids:
                errors.append(f"{ref(item, 'thay_the')}  Văn bản thay thế \"{old}\" không có trong sheet VanBan.")

    # CanCu
    _ws, rows = table("CanCu", CANCU, required=False)
    for item in rows:
        v = item[1]
        if v.get("doc") is None or v.get("van_ban") is None:
            continue
        hit = doc_ids.get(str(v["doc"]))
        if hit is None:
            errors.append(f"{ref(item, 'doc')}  Mã văn bản \"{v['doc']}\" không có trong sheet VanBan.")
            continue
        c = str(v["van_ban"]) if v.get("dieu_khoan") is None else {"van_ban": str(v["van_ban"]),
                                                                    "dieu_khoan": str(v["dieu_khoan"])}
        hit[1].setdefault("can_cu", []).append(c)

    if errors:
        raise ExcelInputError(path, errors)
    out = {"project": project, "attributes": attributes, "steps": steps}
    if contracts:
        out["contracts"] = contracts
    if documents:
        out["documents"] = documents
    return out


# ---------------------------------------------------------------- ghi (mẫu + chuyển YAML)

SUPPORTED_KEYS = {"project", "attributes", "steps", "contracts", "documents"}


def write_workbook(path, rules: dict, cfg: dict, project: dict | None = None) -> None:
    """Tạo file Excel theo mẫu 6 sheet; có `project` thì điền dữ liệu của hồ sơ đó."""
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"{path} đã tồn tại, không ghi đè")
    if project is not None:
        extra = set(project) - SUPPORTED_KEYS
        unknown_attrs = set(project.get("attributes") or {}) - {k for k, _ in ATTRS}
        if extra or unknown_attrs:
            raise ValueError("Hồ sơ có mục mẫu Excel chưa hỗ trợ, không chuyển để tránh mất dữ liệu: "
                             + ", ".join(sorted(extra | unknown_attrs)))
    opx = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = opx.Workbook()
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F0")
    wrap = Alignment(wrap_text=True, vertical="top")

    # HuongDan + danh mục (nguồn cho danh sách chọn - giữ 6 sheet như đặc tả)
    hd = wb.active
    hd.title = "HuongDan"
    lines = [
        "HƯỚNG DẪN ĐIỀN HỒ SƠ DỰ ÁN (cross-check-engine)",
        "Không đọc dữ liệu từ sheet này. Cột có * là bắt buộc; cột có ▼ chọn từ danh sách.",
        "Ngày: ô kiểu ngày, hoặc gõ dd/mm/yyyy, hoặc yyyy-mm-dd.",
        "Số: ô kiểu số; nếu gõ chữ thì dấu chấm phân cách nghìn, dấu phẩy thập phân (120.000.000.000; 24,5).",
        "Ô công thức: engine chỉ đọc giá trị đã tính và lưu trong file.",
        "Sheet CanCu: mỗi dòng một căn cứ viện dẫn của một văn bản; Điều/khoản chỉ ghi khi viện dẫn điều khoản có hiệu lực riêng.",
        "Có lỗi nhập liệu: engine liệt kê toàn bộ lỗi theo Sheet!Ô và không chạy đối soát.",
    ]
    for i, t in enumerate(lines, 1):
        hd.cell(row=i, column=1, value=t).font = bold if i == 1 else Font()
    lists = {"Bước": [r["name"] for r in rules.values()], "Trạng thái": list(STATUS.values()),
             "Nhóm văn bản": list(cfg["groups"].values()), "Có/Không": ["Có", "Không"]}
    start, col, ranges = len(lines) + 2, 1, {}
    hd.cell(row=start, column=1, value="DANH MỤC (dùng cho danh sách chọn)").font = bold
    for name, items in lists.items():
        hd.cell(row=start + 1, column=col, value=name).font = bold
        for j, it in enumerate(items):
            hd.cell(row=start + 2 + j, column=col, value=it)
        letter = opx.utils.get_column_letter(col)
        ranges[name] = f"'HuongDan'!${letter}${start + 2}:${letter}${start + 1 + len(items)}"
        hd.column_dimensions[letter].width = 48 if name in ("Bước", "Nhóm văn bản") else 18
        col += 1

    def dv(ws, list_name, col_letter, rows=500):
        v = DataValidation(type="list", formula1=ranges[list_name], allow_blank=True, showErrorMessage=True,
                           errorTitle="Giá trị không hợp lệ", error=f"Chọn trong danh mục {list_name} (sheet HuongDan).")
        ws.add_data_validation(v)
        v.add(f"{col_letter}2:{col_letter}{rows}")

    def header(ws, cols):
        for c, (k, label) in enumerate(cols, 1):
            mark = " ▼" if k in ("step", "status", "nhom") else ""
            cell = ws.cell(row=1, column=c, value=label + mark)
            cell.font, cell.fill, cell.alignment = bold, head_fill, wrap
            ws.column_dimensions[opx.utils.get_column_letter(c)].width = WIDE.get(k, max(14, min(40, len(label) + 4)))
        ws.freeze_panes = "A2"

    def put(ws, row, col, v, is_date=False, money=False):
        cell = ws.cell(row=row, column=col, value=v)
        if is_date and v is not None:
            cell.number_format = "dd/mm/yyyy"
        if money:
            cell.number_format = "#,##0"  # chỉ đổi hiển thị (phân cách nghìn), giá trị ô giữ nguyên

    # DuAn
    da = wb.create_sheet("DuAn")
    for c, label in enumerate(("Mục", "Giá trị"), 1):
        cell = da.cell(row=1, column=c, value=label)
        cell.font, cell.fill = bold, head_fill
    da.column_dimensions["A"].width, da.column_dimensions["B"].width = 36, 40
    meta, attrs = (project or {}).get("project") or {}, (project or {}).get("attributes") or {}
    da_rows = [("Mã dự án *", meta.get("id")), ("Tên dự án *", meta.get("name"))]
    da_rows += [(label + " ▼", None if k not in attrs else ("Có" if attrs[k] else "Không")) for k, label in ATTRS]
    for i, (label, v) in enumerate(da_rows, 2):
        da.cell(row=i, column=1, value=label)
        da.cell(row=i, column=2, value=v)
    yn = DataValidation(type="list", formula1=ranges["Có/Không"], allow_blank=True)
    da.add_data_validation(yn)
    yn.add(f"B4:B{len(da_rows) + 1}")

    # TrinhTu
    tt = wb.create_sheet("TrinhTu")
    header(tt, TRINHTU)
    dv(tt, "Bước", "A")
    dv(tt, "Trạng thái", "B")
    for i, (sid, rec) in enumerate(((project or {}).get("steps") or {}).items(), 2):
        rec = rec or {}
        put(tt, i, 1, rules[sid]["name"] if sid in rules else sid)
        put(tt, i, 2, STATUS.get(rec.get("status"), rec.get("status")))
        put(tt, i, 3, rec.get("start_date"), True)
        put(tt, i, 4, rec.get("date"), True)
        put(tt, i, 5, rec.get("evidence"))
        put(tt, i, 6, rec.get("na_basis"))

    # HopDong
    hdg = wb.create_sheet("HopDong")
    header(hdg, HOPDONG)
    for i, (cid, c) in enumerate(((project or {}).get("contracts") or {}).items(), 2):
        put(hdg, i, 1, cid)
        put(hdg, i, 2, (c or {}).get("ngay_ky"), True)

    # VanBan + CanCu
    vb, cc = wb.create_sheet("VanBan"), wb.create_sheet("CanCu")
    header(vb, VANBAN)
    header(cc, CANCU)
    dv(vb, "Nhóm văn bản", "C")
    for c, (k, _label) in enumerate(VANBAN, 1):  # ô trống của mẫu cũng có sẵn định dạng ngày / tiền
        if k in DATE_COLS or k in MONEY_COLS:
            for r in range(2, 201):
                vb.cell(row=r, column=c).number_format = "dd/mm/yyyy" if k in DATE_COLS else "#,##0"
    crow = 2
    for i, d in enumerate((project or {}).get("documents") or [], 2):
        info = d.get("thong_tin") or {}
        tg = info.get("thoi_gian_thuc_hien") or {}
        vals = {"id": d.get("id"), "loai": d.get("loai"), "nhom": cfg["groups"].get(d.get("nhom"), d.get("nhom")),
                "ngay_ky": d.get("ngay_ky"), "ngay_nop": d.get("ngay_nop"),
                "ngay_phat_hanh_hsmt": d.get("ngay_phat_hanh_hsmt"), "hop_dong": d.get("hop_dong"),
                "thay_the": ", ".join(d.get("thay_the") or []) or None,
                **{k: info.get(k) for k in INFO_TEXT + INFO_NUM}, "tg_tu": tg.get("tu"), "tg_den": tg.get("den")}
        for c, (k, _label) in enumerate(VANBAN, 1):
            put(vb, i, c, vals[k], k in DATE_COLS, k in MONEY_COLS)
        for cit in d.get("can_cu") or []:
            put(cc, crow, 1, d.get("id"))
            if isinstance(cit, dict):
                put(cc, crow, 2, cit.get("van_ban"))
                put(cc, crow, 3, cit.get("dieu_khoan"))
            else:
                put(cc, crow, 2, cit)
            crow += 1
    wb.save(path)


def main(argv=None) -> int:
    import yaml

    from .cli import DEFAULT_CONSISTENCY, DEFAULT_RULES
    from .consistency import load_consistency_rules
    from .engine import load_rules

    args = list(sys.argv[1:] if argv is None else argv)
    usage = ("Dùng: python3 -m crosscheck.excel mau <ra.xlsx>\n"
             "      python3 -m crosscheck.excel tu-yaml <ho_so.yaml> <ra.xlsx>")
    if not args or args[0] not in ("mau", "tu-yaml") or len(args) != (2 if args[0] == "mau" else 3):
        print(usage, file=sys.stderr)
        return 2
    try:
        rules, cfg = load_rules(DEFAULT_RULES), load_consistency_rules(DEFAULT_CONSISTENCY)
        if args[0] == "mau":
            write_workbook(args[1], rules, cfg)
            print(f"Đã tạo file mẫu: {args[1]}")
        else:
            with open(args[1], encoding="utf-8") as f:
                project = yaml.safe_load(f)
            write_workbook(args[2], rules, cfg, project)
            print(f"Đã chuyển {args[1]} → {args[2]}")
    except (FileExistsError, ValueError, OSError, RuleError) as e:
        print(f"Lỗi: {e}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
